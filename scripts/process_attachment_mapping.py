#!/usr/bin/env python3
import argparse
import os
import re
from copy import copy
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import openpyxl
from openpyxl.utils import get_column_letter

SPECIAL_COLS = {
    "10001": ["射程"],
    "10005": ["后坐力控制", "后坐"],
    "100005": ["后坐力控制", "后坐"],
    "10006": ["操控速度", "操控"],
    "10007": ["据枪稳定性", "稳定性"],
    "10008": ["腰射精度", "腰射", "腰际射击精度"],
}
ERROR_VALUES = {"#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?"}
RESULT_HEADERS = ["映射后优点", "映射后缺点", "优点校验"]


def norm_id(v):
    if v is None:
        return ""
    s = str(v).strip()
    return s[:-2] if s.endswith(".0") else s


def fmt_num(v, percent=False):
    if v is None or v == "":
        return ""
    try:
        x = float(v)
    except Exception:
        return str(v)
    if percent:
        x *= 100
        s = f"{x:.10g}"
        if "." in s:
            s = s.rstrip("0").rstrip(".")
        return ("+" if x > 0 else "") + s + "%"
    if abs(x - round(x)) < 1e-9:
        x = int(round(x))
    return ("+" if x > 0 else "") + str(x)


def output_default(input_path):
    base_name = os.path.basename(input_path)
    base, ext = os.path.splitext(base_name)
    return os.path.abspath(f"{base}_文本映射版{ext}")


def repair_empty_fills(input_path, repair_dir):
    src = Path(input_path)
    out = Path(repair_dir) / f"{src.stem}_样式修复中间版{src.suffix}"
    with ZipFile(src, "r") as zin, ZipFile(out, "w", ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/styles.xml":
                text = data.decode("utf-8", errors="replace")
                text = text.replace('<fill/>', '<fill><patternFill patternType="none"/></fill>')
                data = text.encode("utf-8")
            zout.writestr(item, data)
    return str(out)


def load_workbook_safe(input_path, repair_dir):
    try:
        return openpyxl.load_workbook(input_path, data_only=False), None
    except TypeError as e:
        if "Fill" not in str(e):
            raise
        repaired = repair_empty_fills(input_path, repair_dir)
        return openpyxl.load_workbook(repaired, data_only=False), repaired


def headers_of(ws):
    return {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1) if ws.cell(1, c).value is not None}


def find_header(ws, candidates, fallback=None):
    headers = headers_of(ws)
    for name in candidates:
        if name in headers:
            return headers[name]
    for header, col in headers.items():
        header_s = str(header)
        for name in candidates:
            if name in header_s or header_s in name:
                return col
    if fallback is not None:
        return fallback
    raise ValueError(f"找不到字段：{candidates}；当前表头={list(headers.keys())}")


def find_attr_col(headers, candidates):
    for name in candidates:
        if name in headers:
            return headers[name]
    for header, col in headers.items():
        header_s = str(header)
        for name in candidates:
            if name in header_s or header_s in name:
                return col
    return None


def looks_like_buff_sheet(ws):
    headers = headers_of(ws)
    has_adv = any("优点" in str(h) for h in headers)
    has_dis = any("缺点" in str(h) for h in headers)
    if has_adv and has_dis:
        return True
    max_r = min(ws.max_row, 30)
    max_c = min(ws.max_column, 80)
    for row in ws.iter_rows(min_row=1, max_row=max_r, max_col=max_c):
        for cell in row:
            if isinstance(cell.value, str) and ("BuffText" in cell.value or "Params=" in cell.value):
                return True
    return False


def looks_like_map_sheet(ws):
    headers = headers_of(ws)
    header_text = " ".join(str(h) for h in headers)
    if "ID" in header_text and ("描述" in header_text or "不需要翻译" in header_text):
        return True
    name_hit = "映射" in ws.title and ("描述" in ws.title or "语言" in ws.title)
    first = norm_id(ws.cell(1, 1).value)
    second = str(ws.cell(1, 2).value or "")
    return name_hit or (first.upper() == "ID" and ("描述" in second or "value" in second))


def detect_sheets(wb, data_sheet=None, map_sheet=None):
    if data_sheet and data_sheet not in wb.sheetnames:
        raise ValueError(f"找不到数据 sheet：{data_sheet}；当前 sheets={wb.sheetnames}")
    if map_sheet and map_sheet not in wb.sheetnames:
        raise ValueError(f"找不到映射 sheet：{map_sheet}；当前 sheets={wb.sheetnames}")

    if not data_sheet:
        candidates = [ws.title for ws in wb.worksheets if looks_like_buff_sheet(ws)]
        if not candidates:
            raise ValueError("未自动识别到包含 BuffText/Params 或优缺点字段的数据 sheet")
        data_sheet = sorted(candidates, key=lambda x: ("配件" not in x, x))[0]

    if not map_sheet:
        candidates = [ws.title for ws in wb.worksheets if looks_like_map_sheet(ws)]
        if not candidates:
            raise ValueError("未自动识别到映射关系 sheet")
        map_sheet = sorted(candidates, key=lambda x: ("配件" not in x, "MA2" not in x, x))[0]

    return data_sheet, map_sheet


def build_mapping(map_ws):
    id_col = find_header(map_ws, ["ID"], fallback=1)
    desc_col = find_header(map_ws, ["描述（{value1}不需要翻译）", "描述", "文本"], fallback=2)
    mapping = {}
    for r in range(2, map_ws.max_row + 1):
        key = norm_id(map_ws.cell(r, id_col).value)
        if key:
            mapping[key] = map_ws.cell(r, desc_col).value
    return mapping


def replace_placeholders(text, params):
    vals = [p for p in params if p != ""]
    for i, val in enumerate(vals, 1):
        text = text.replace(f"{{value{i}}}", val)
    return text


def render(raw, row_idx, ws, headers, mapping, missing, missing_attr):
    if not raw:
        return ""
    out = []
    for params_str in re.findall(r"Params=\(([^)]*)\)", str(raw)):
        params = re.findall(r'"([^"]*)"', params_str)
        if not params:
            continue
        code = norm_id(params[-1])
        text = mapping.get(code)
        if text is None:
            missing.add(code)
            out.append(f"未匹配:{code}")
            continue
        text = str(text)
        if code in SPECIAL_COLS:
            col_idx = find_attr_col(headers, SPECIAL_COLS[code])
            if col_idx is None:
                missing_attr.add(f"{code}:{'/'.join(SPECIAL_COLS[code])}")
                out.append(text)
            else:
                value = ws.cell(row_idx, col_idx).value
                out.append(f"{text}{fmt_num(value, code == '10001')}")
        else:
            out.append(replace_placeholders(text, params[:-1]))
    return "；".join(out)


def copy_style(src, dst):
    try:
        if src.has_style:
            dst._style = copy(src._style)
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.alignment = copy(src.alignment)
        dst.border = copy(src.border)
        dst.number_format = src.number_format
    except Exception:
        pass


def scan_errors(wb):
    errors = []
    formulas = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas += 1
                if cell.value in ERROR_VALUES:
                    errors.append(f"{ws.title}!{cell.coordinate}:{cell.value}")
    return formulas, errors


def process(input_path, output_path=None, data_sheet=None, map_sheet=None):
    if output_path is None:
        output_path = output_default(input_path)
    output_path = os.path.abspath(output_path)
    wb, repaired = load_workbook_safe(input_path, Path(output_path).parent)
    data_sheet, map_sheet = detect_sheets(wb, data_sheet, map_sheet)

    ws = wb[data_sheet]
    mapping = build_mapping(wb[map_sheet])
    headers = headers_of(ws)

    adv_col = find_header(ws, ["优点（文本+电池版）", "优点", "优点原文"])
    dis_col = find_header(ws, ["缺点（文本+电池版）", "缺点", "缺点原文"])
    actual_adv_col = headers.get("实际优点")

    start_col = ws.max_column + 1
    for h in RESULT_HEADERS:
        if h in headers:
            start_col = min(start_col, headers[h])
    new_cols = [start_col, start_col + 1, start_col + 2]

    style_src = ws.cell(1, actual_adv_col or adv_col)
    for col, title in zip(new_cols, RESULT_HEADERS):
        cell = ws.cell(1, col)
        cell.value = title
        copy_style(style_src, cell)

    missing = set()
    missing_attr = set()
    processed = 0
    matched = 0
    mismatched = []
    for r in range(2, ws.max_row + 1):
        raw_adv = ws.cell(r, adv_col).value
        raw_dis = ws.cell(r, dis_col).value
        if not raw_adv and not raw_dis:
            continue
        processed += 1
        adv = render(raw_adv, r, ws, headers, mapping, missing, missing_attr)
        dis = render(raw_dis, r, ws, headers, mapping, missing, missing_attr)
        ws.cell(r, new_cols[0]).value = adv or None
        ws.cell(r, new_cols[1]).value = dis or None
        if actual_adv_col and ws.cell(r, actual_adv_col).value and adv:
            actual = str(ws.cell(r, actual_adv_col).value)
            ok = actual == adv
            ws.cell(r, new_cols[2]).value = "一致" if ok else "不一致"
            matched += 1 if ok else 0
            if not ok:
                mismatched.append(r)
        for col in new_cols:
            copy_style(ws.cell(r, actual_adv_col or adv_col), ws.cell(r, col))

    for col, width in zip(new_cols, [45, 32, 12]):
        ws.column_dimensions[get_column_letter(col)].width = width

    wb.save(output_path)
    formulas, errors = scan_errors(wb)
    return {
        "output": output_path,
        "data_sheet": data_sheet,
        "map_sheet": map_sheet,
        "repaired_intermediate": repaired,
        "processed_rows": processed,
        "matched_actual_advantage": matched,
        "mismatched_rows": mismatched[:30],
        "missing_codes": sorted(missing),
        "missing_attr_columns": sorted(missing_attr),
        "formula_count": formulas,
        "error_count": len(errors),
        "errors": errors[:20],
    }


def main():
    parser = argparse.ArgumentParser(description="批量生成三角洲行动配件文本版优缺点")
    parser.add_argument("input", help="输入 Excel 路径")
    parser.add_argument("-o", "--output", help="输出 Excel 路径；不填则在当前目录生成 _文本映射版")
    parser.add_argument("--data-sheet", default=None, help="数据 sheet 名；不填自动识别")
    parser.add_argument("--map-sheet", default=None, help="映射关系 sheet 名；不填自动识别")
    args = parser.parse_args()
    result = process(args.input, args.output, args.data_sheet, args.map_sheet)
    print("处理完成")
    for k, v in result.items():
        print(f"{k}: {v}")


if __name__ == "__main__":
    main()
